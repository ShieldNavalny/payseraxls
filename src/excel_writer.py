from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from pathlib import Path
from typing import List, Dict

class ExcelWriter:
    def __init__(self):
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.worksheet.title = "Проблемные транзакции"
        
        self.red_fill = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
        self.purple_fill = PatternFill(start_color="FF9933FF", end_color="FF9933FF", fill_type="solid")
        self.white_font = Font(color="FFFFFFFF", bold=True, size=11)
        self.border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
    
    def write_transactions(self, transactions: List[Dict], output_path: str):
        headers = ['Дата', 'Получатель', 'Сумма', 'Карта', 'Комментарий']
        
        for col_num, header in enumerate(headers, 1):
            cell = self.worksheet.cell(row=1, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFFFF", size=12)
            cell.fill = PatternFill(start_color="FF2C3E50", end_color="FF2C3E50", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = self.border
        
        for row_num, t in enumerate(transactions, 2):
            # Отрицательная сумма → красный, положительная → фиолетовый
            is_negative = t['amount'].startswith('-')
            fill = self.red_fill if is_negative else self.purple_fill
            
            row_data = [
                t['date'],
                t['recipient'],
                t['amount'],
                t['card_last_four'] if t['card_last_four'] != 'N/A' else '',
                t['comment']
            ]
            
            for col_num, value in enumerate(row_data, 1):
                cell = self.worksheet.cell(row=row_num, column=col_num)
                cell.value = value
                cell.fill = fill
                cell.font = self.white_font
                cell.border = self.border
                cell.alignment = Alignment(
                    horizontal="left" if col_num in [2, 5] else "center",
                    vertical="center"
                )
        
        self.worksheet.column_dimensions['A'].width = 13
        self.worksheet.column_dimensions['B'].width = 40
        self.worksheet.column_dimensions['C'].width = 15
        self.worksheet.column_dimensions['D'].width = 10
        self.worksheet.column_dimensions['E'].width = 30
        
        self.worksheet.freeze_panes = "A2"
        
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        self.workbook.save(output_path)
        print(f"✓ Отчет: {output_path}")
