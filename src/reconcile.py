"""
reconcile.py — сравнение XLS-вывода с CSV-реестром счетов.

Логика:
  1. Читает CSV из корня рабочей директории (documents.csv по умолчанию).
  2. Читает XLSX из output_reports/.
  3. Строит ключи (дата, canonical_vendor, |сумма|) для обеих сторон.
     - XLS «Получатель» нормализуется через VENDOR_MAP.
     - CSV «Vendor» тоже нормализуется через тот же VENDOR_MAP.
     - Суммы в CSV всегда положительные; в XLS могут быть отрицательными — берём |abs|.
  4. Находит записи из CSV, которых НЕТ в XLS (только расходы — отрицательные в XLS).
  5. Копирует недостающие PDF в missing_invoices/<Vendor>/
"""

import csv
import re
import shutil
from pathlib import Path
from typing import List, Dict, Set, Tuple, Optional

import openpyxl

from config import OUTPUT_DIR, OUTPUT_FILENAME

# ─── настройки ────────────────────────────────────────────────────────────────
DEFAULT_CSV_NAME = "documents.csv"
MISSING_DIR_NAME = "missing_invoices"

# ─── словарь нормализации вендоров ────────────────────────────────────────────
# Ключи — подстроки (lowercase), встречающиеся в поле «Получатель» XLS
#          ИЛИ в поле «Vendor» CSV.
# Значение — канонический идентификатор вендора.
#
# Правило: берём первое совпадение (более специфичные — выше в списке).
# Добавляй новые строки по мере появления новых контрагентов.
VENDOR_MAP: List[Tuple[str, str]] = [
    # XLS-строка              → canonical
    ("sinch mailgun",         "mailgun"),
    ("mailgun",               "mailgun"),
    ("hetzner online",        "hetzner"),
    ("hetzner",               "hetzner"),
    ("amazon",                "amazon"),
    ("allegro",               "allegro"),
    ("aliexpress",            "aliexpress"),
    ("bolt.eu",               "bolt"),
    ("bolt",                  "bolt"),
    ("eu.store.ui.com",       "ui.com"),
    ("ui.com",                "ui.com"),
    ("eurocash",              "eurocash1.lt"),
    ("pirkeu",                "pirkeu.lt"),
    ("sandeliukunuoma",       "sandeliukunuoma.lt"),
    # добавляй сюда новые пары по мере необходимости:
    # ("partial name in xls", "canonical"),
]


def _canonical_vendor(raw: str) -> str:
    """
    Возвращает канонический идентификатор вендора или lowercase оригинал,
    если совпадение не найдено.
    """
    low = raw.strip().lower()
    for pattern, canonical in VENDOR_MAP:
        if pattern in low:
            return canonical
    return low  # fallback — без нормализации


# ─── нормализация суммы ───────────────────────────────────────────────────────

def _normalize_amount(raw: str) -> str:
    """
    Приводит сумму к строке без знака и EUR, с двумя знаками после точки.
    Работает и для '-32.00 EUR', и для '32.00', и для '32'.
    """
    cleaned = raw.strip().replace(' EUR', '').replace(',', '.').replace('\xa0', '')
    try:
        return f"{abs(float(cleaned)):.2f}"
    except ValueError:
        return cleaned


# ─── нормализация даты ────────────────────────────────────────────────────────

def _normalize_date(raw) -> str:
    """
    Принимает datetime-объект, строку '02.01.2025' или '2025-01-02' →
    всегда возвращает 'YYYY-MM-DD'.
    """
    if hasattr(raw, 'strftime'):
        return raw.strftime('%Y-%m-%d')
    raw = str(raw).strip()
    m = re.match(r'^(\d{2})\.(\d{2})\.(\d{4})$', raw)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    return raw


# ─── ключ для матчинга ────────────────────────────────────────────────────────

def _make_key(date: str, vendor: str, amount: str) -> Tuple[str, str, str]:
    return (
        _normalize_date(date),
        _canonical_vendor(vendor),
        _normalize_amount(amount),
    )


# ─── чтение XLS ───────────────────────────────────────────────────────────────

def load_xls_keys(xls_path: Path) -> Set[Tuple[str, str, str]]:
    """
    Читает все листы XLS (кроме 'Сводка') и собирает ключи расходных транзакций.
    Колонки листа: Дата(1) Получатель(2) Сумма(3) ...

    Важно: мы сверяем ТОЛЬКО расходы (сумма < 0 в XLS), потому что
    в CSV хранятся только счета/инвойсы (исходящие платежи).
    Возвраты (положительные строки XLS) в матчинг не включаем.
    """
    keys: Set[Tuple[str, str, str]] = set()
    wb = openpyxl.load_workbook(xls_path, data_only=True)

    for sheet_name in wb.sheetnames:
        if sheet_name == 'Сводка':
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2, values_only=True):
            date_val, vendor_val, amount_val = row[0], row[1], row[2]
            if not date_val or amount_val is None:
                continue

            # Пропускаем строки итогов (нет даты или получателя)
            if not vendor_val:
                continue

            amount_str = str(amount_val).replace(' EUR', '').replace(',', '.').strip()
            try:
                amount_float = float(amount_str)
            except ValueError:
                continue

            # Берём только расходы (отрицательные суммы)
            if amount_float >= 0:
                continue

            keys.add(_make_key(
                str(date_val),
                str(vendor_val),
                str(amount_val),
            ))

    wb.close()
    return keys


# ─── чтение CSV ───────────────────────────────────────────────────────────────

def load_csv_records(csv_path: Path) -> List[Dict]:
    """
    Возвращает список записей из CSV.
    Ожидаемые колонки: Date, Vendor, Amount, Currency, Invoice ID, Type, File

    Суммы в CSV всегда положительные (абсолютные значения расходов).
    """
    records = []
    with open(csv_path, newline='', encoding='utf-8-sig') as fh:
        reader = csv.DictReader(fh)
        for row in reader:
            records.append({
                'date':       row.get('Date', '').strip(),
                'vendor':     row.get('Vendor', '').strip(),
                'amount':     row.get('Amount', '').strip(),
                'invoice_id': row.get('Invoice ID', '').strip(),
                'type':       row.get('Type', '').strip(),
                'file':       row.get('File', '').strip(),
            })
    return records


# ─── основная функция ─────────────────────────────────────────────────────────

def reconcile(
    base_dir: Optional[Path] = None,
    csv_name: str = DEFAULT_CSV_NAME,
    xls_name: Optional[str] = None,
    missing_dir_name: str = MISSING_DIR_NAME,
    dry_run: bool = False,
) -> List[Dict]:
    """
    Запускает сверку.

    :param base_dir:         корневая директория (по умолчанию — cwd).
    :param csv_name:         имя CSV-файла в base_dir.
    :param xls_name:         имя XLS-файла (по умолчанию из config.py).
    :param missing_dir_name: имя выходной директории.
    :param dry_run:          если True — только вычисляет, не копирует.
    :return:                 список недостающих записей.
    """
    if base_dir is None:
        base_dir = Path.cwd()

    csv_path = base_dir / csv_name
    if not csv_path.exists():
        raise FileNotFoundError(f"CSV не найден: {csv_path}")

    xls_path = OUTPUT_DIR / (xls_name or OUTPUT_FILENAME)
    if not xls_path.exists():
        raise FileNotFoundError(f"XLS не найден: {xls_path}")

    print(f"📂 CSV:  {csv_path}")
    print(f"📊 XLS:  {xls_path}")

    xls_keys = load_xls_keys(xls_path)
    print(f"   XLS расходных транзакций: {len(xls_keys)}")

    csv_records = load_csv_records(csv_path)
    print(f"   CSV записей:              {len(csv_records)}")

    # ── находим недостающие ───────────────────────────────────────────────────
    missing: List[Dict] = []
    for rec in csv_records:
        key = _make_key(rec['date'], rec['vendor'], rec['amount'])
        if key not in xls_keys:
            missing.append(rec)

    print(f"\n🔍 Недостающих счетов: {len(missing)}")

    if not missing:
        print("✓ Всё совпадает — расхождений нет.")
        return missing

    # ── группируем по вендору для наглядности ─────────────────────────────────
    by_vendor: Dict[str, List[Dict]] = {}
    for rec in missing:
        by_vendor.setdefault(rec['vendor'], []).append(rec)

    for vendor, recs in sorted(by_vendor.items()):
        print(f"   {vendor}: {len(recs)} шт.")

    if dry_run:
        print("\n⚠  Режим dry-run: файлы не скопированы.")
        _print_missing_table(missing)
        return missing

    # ── копируем PDF ──────────────────────────────────────────────────────────
    out_root = base_dir / missing_dir_name
    copied = 0
    not_found = 0

    print(f"\n📁 Копируем в: {out_root}")

    for rec in missing:
        if not rec['file']:
            print(f"   ⚠ Нет пути к файлу для: {rec['vendor']} {rec['date']} {rec['amount']}")
            continue

        src = base_dir / rec['file']
        if not src.exists():
            print(f"   ⚠ Не найден файл: {src}")
            not_found += 1
            continue

        # Папка: missing_invoices/<Vendor>/
        vendor_safe = re.sub(r'[\\/:*?"<>|]', '_', rec['vendor'])
        dst_dir = out_root / vendor_safe
        dst_dir.mkdir(parents=True, exist_ok=True)

        dst = dst_dir / src.name
        shutil.copy2(src, dst)
        copied += 1

    print(f"\n✓ Скопировано: {copied}")
    if not_found:
        print(f"⚠  Не найдено на диске: {not_found}")

    return missing


def _print_missing_table(records: List[Dict]) -> None:
    """Печатает таблицу недостающих записей (для dry-run)."""
    print(f"\n{'Дата':<12} {'Вендор':<25} {'Сумма':>10}  Invoice ID")
    print("-" * 70)
    for r in records:
        print(f"{r['date']:<12} {r['vendor']:<25} {r['amount']:>10}  {r['invoice_id']}")


# ─── CLI ──────────────────────────────────────────────────────────────────────

def main():
    import argparse

    parser = argparse.ArgumentParser(
        description="Сверка XLS-вывода с CSV-реестром счетов"
    )
    parser.add_argument(
        '--csv', default=DEFAULT_CSV_NAME,
        help=f"Имя CSV-файла в корне (по умолчанию: {DEFAULT_CSV_NAME})"
    )
    parser.add_argument(
        '--xls', default=None,
        help="Имя XLS-файла в output_reports/ (по умолчанию из config.py)"
    )
    parser.add_argument(
        '--out', default=MISSING_DIR_NAME,
        help=f"Папка для недостающих счетов (по умолчанию: {MISSING_DIR_NAME})"
    )
    parser.add_argument(
        '--dry-run', action='store_true',
        help="Только показать расхождения, не копировать файлы"
    )
    parser.add_argument(
        '--base-dir', default=None,
        help="Корневая директория (по умолчанию: текущая)"
    )
    args = parser.parse_args()

    base = Path(args.base_dir) if args.base_dir else None

    print("🔄 Payseraxls — Reconcile")
    print("=" * 50)
    reconcile(
        base_dir=base,
        csv_name=args.csv,
        xls_name=args.xls,
        missing_dir_name=args.out,
        dry_run=args.dry_run,
    )
    print("=" * 50)


if __name__ == '__main__':
    main()
